"""Bounded extraction subprocess. Output is data, never executable instructions."""
import json
import sys
import zipfile
from pathlib import Path

MAX_TEXT = 5_000_000
MAX_PAGES = 2000


def extract(path: Path, extension: str, ranges=None):
    sections = []
    count = 0
    if extension == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError('PDF terkunci; unggah salinan yang dapat dibaca dan diizinkan.')
        if len(reader.pages) > MAX_PAGES:
            raise ValueError('PDF melebihi batas 2.000 halaman. File asli tetap disimpan.')
        total_pages = len(reader.pages)
        if ranges and any(r['end'] > total_pages for r in ranges):
            raise ValueError('Rentang pilihan melebihi jumlah halaman PDF; perbaiki metadata.')
        selected = sorted({n for r in ranges for n in range(r['start'], r['end'] + 1)}) if ranges else range(1, total_pages + 1)
        for number in selected:
            page = reader.pages[number - 1]
            text = page.extract_text() or ''
            count += len(text)
            if count > MAX_TEXT:
                raise ValueError('Teks melebihi batas pilot; perlu pemrosesan dokumen besar.')
            sections.append({'locator': f'Halaman PDF {number}', 'page': number, 'text': text})
    else:
        from docx import Document
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 80 * 1024 * 1024 or len(archive.infolist()) > 5000:
                raise ValueError('DOCX melewati batas ukuran dekompresi.')
            if 'word/document.xml' not in archive.namelist():
                raise ValueError('Berkas bukan DOCX yang valid.')
            if any('vbaProject' in i for i in archive.namelist()):
                raise ValueError('Dokumen bermakro tidak didukung.')
        doc = Document(path)
        for number, paragraph in enumerate(doc.paragraphs, 1):
            if paragraph.text.strip():
                sections.append({'locator': f'Paragraf {number}', 'text': paragraph.text})
        for number, table in enumerate(doc.tables, 1):
            sections.append({'locator': f'Tabel {number}', 'text': '\n'.join(' | '.join(c.text for c in row.cells) for row in table.rows)})
        count = sum(len(s['text']) for s in sections)
        if count > MAX_TEXT:
            raise ValueError('Teks melebihi batas pilot.')
    usable = sum(len(s['text'].strip()) for s in sections)
    return {'status': 'readable' if usable >= 100 else 'needs_review', 'characters': usable,
            'sections': sections, 'message': '' if usable >= 100 else 'Teks terlalu sedikit; periksa kebutuhan OCR atau isi dokumen.',
            'method': 'pdf-text' if extension == '.pdf' else 'docx-paragraphs-tables',
            'total_pages': total_pages if extension == '.pdf' else None,
            'selected_pages': len(sections) if extension == '.pdf' else None}


if __name__ == '__main__':
    try:
        if sys.platform != 'win32':
            import resource
            resource.setrlimit(resource.RLIMIT_AS, (1024 * 1024 * 1024, 1024 * 1024 * 1024))
            resource.setrlimit(resource.RLIMIT_CPU, (170, 175))
        result = extract(Path(sys.argv[1]), sys.argv[2], json.loads(sys.argv[3]) if len(sys.argv) > 3 else None)
    except ValueError as exc:
        result = {'status': 'failed', 'characters': 0, 'sections': [], 'message': str(exc)}
    except Exception:
        result = {'status': 'failed', 'characters': 0, 'sections': [], 'message': 'Ekstraksi gagal. File mungkin rusak atau formatnya tidak didukung.'}
    print(json.dumps(result, ensure_ascii=True))
